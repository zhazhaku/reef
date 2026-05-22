"""
P5.3: End-to-end integration test for the Reef PPT Agent.
Full 5-phase pipeline simulation.
"""
import unittest
import json
import os
import sys
import shutil

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "pkg", "skills", "ppt-agent"))

from template_parser import parse_template
from ppt_compositor import compose
from schema import validate_template_structure, validate_style_mapping

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
TEMPLATE_PATH = os.path.join(FIXTURES_DIR, "test_template.pptx")
E2E_DIR = os.path.join(FIXTURES_DIR, "e2e_output")


class TestE2EPPTAgent(unittest.TestCase):
    """P5.3: Full pipeline end-to-end test."""

    @classmethod
    def setUpClass(cls):
        if os.path.exists(E2E_DIR):
            shutil.rmtree(E2E_DIR)
        os.makedirs(E2E_DIR, exist_ok=True)
        cls.workspace = E2E_DIR

    # ── Phase 1: Template Parse ─────────────────────────────────────────────

    def test_e2e_phase1_parse(self):
        """Phase 1: Parse template → valid JSON output."""
        result = parse_template(TEMPLATE_PATH)
        errors = validate_template_structure(result)
        self.assertEqual(errors, [], f"Template validation errors: {errors}")

        structure_path = os.path.join(self.workspace, "template_structure.json")
        with open(structure_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        self.assertTrue(os.path.exists(structure_path))
        self.assertEqual(len(result["slides"]), 5)

    # ── Phase 2-5: Full pipeline ─────────────────────────────────────────────

    def test_e2e_full_pipeline(self):
        """P5.3.1-P5.3.5: Full 5-phase pipeline → verified .pptx output."""
        from pptx import Presentation
        from pptx.util import Pt

        # Phase 1: Parse
        result = parse_template(TEMPLATE_PATH)
        structure_path = os.path.join(self.workspace, "template_structure.json")
        with open(structure_path, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)

        # Phase 2: Outline (manual for test)
        outline = {
            "title": "Project Reef",
            "slides": [
                {"slide_number": 1, "title": "Project Reef", "summary": "Introduction", "slide_type": "COVER"},
                {"slide_number": 2, "title": "Features", "summary": "Capabilities", "slide_type": "CONTENT"},
                {"slide_number": 3, "title": "Architecture", "summary": "Design", "slide_type": "CONTENT"},
                {"slide_number": 4, "title": "Next Steps", "summary": "Future", "slide_type": "CONTENT"},
                {"slide_number": 5, "title": "Thank You", "summary": "Close", "slide_type": "ENDING"},
            ]
        }
        outline_path = os.path.join(self.workspace, "outline.json")
        with open(outline_path, "w", encoding="utf-8") as f:
            json.dump(outline, f, ensure_ascii=False, indent=2)

        # Phase 3: Detail plan (manual for test)
        plan = {
            "title": "Project Reef",
            "slides": [
                {"slide_number": 1, "slide_type": "COVER", "title": "Project Reef",
                 "content_blocks": [
                     {"type": "title", "content": "Project Reef", "overrides": {"font_size": 44}},
                     {"type": "subtitle", "content": "Distributed AI Swarm", "overrides": {}},
                 ]},
                {"slide_number": 2, "slide_type": "CONTENT", "title": "Features",
                 "content_blocks": [
                     {"type": "title", "content": "Key Features", "overrides": {}},
                     {"type": "bullet_list", "items": ["Multi-agent", "Skill routing", "Sandboxed exec"],
                      "overrides": {"font_size": 18}},
                 ]},
                {"slide_number": 3, "slide_type": "CONTENT", "title": "Architecture",
                 "content_blocks": [
                     {"type": "title", "content": "Architecture", "overrides": {}},
                     {"type": "text", "content": "Distributed system with central server and agent clients.", "overrides": {}},
                 ]},
                {"slide_number": 4, "slide_type": "CONTENT", "title": "Next Steps",
                 "content_blocks": [
                     {"type": "title", "content": "Next Steps", "overrides": {}},
                     {"type": "bullet_list", "items": ["Deploy", "Add roles", "Dashboard"],
                      "overrides": {"font_size": 18}},
                 ]},
                {"slide_number": 5, "slide_type": "ENDING", "title": "Thank You",
                 "content_blocks": [
                     {"type": "title", "content": "Thank You", "overrides": {"font_size": 48, "color": "#1A56DB"}},
                     {"type": "subtitle", "content": "Q & A", "overrides": {}},
                 ]},
            ]
        }
        detail_path = os.path.join(self.workspace, "detail_plan.json")
        with open(detail_path, "w", encoding="utf-8") as f:
            json.dump(plan, f, ensure_ascii=False, indent=2)

        # Phase 4: Style mapping
        mapping = {"mappings": [
            {"plan_slide": 1, "template_slide": 1, "reason": "Cover"},
            {"plan_slide": 2, "template_slide": 2, "reason": "Content bullets"},
            {"plan_slide": 3, "template_slide": 2, "reason": "Content text"},
            {"plan_slide": 4, "template_slide": 2, "reason": "Content bullets"},
            {"plan_slide": 5, "template_slide": 5, "reason": "Ending"},
        ]}
        mapping_path = os.path.join(self.workspace, "style_mapping.json")
        with open(mapping_path, "w", encoding="utf-8") as f:
            json.dump(mapping, f, ensure_ascii=False, indent=2)
        errors = validate_style_mapping(mapping, 5)
        self.assertEqual(errors, [])

        # Phase 5: Compose
        images_dir = os.path.join(self.workspace, "images")
        os.makedirs(images_dir, exist_ok=True)
        output_path = os.path.join(self.workspace, "result.pptx")

        compose(TEMPLATE_PATH, detail_path, mapping_path, images_dir, output_path)
        self.assertTrue(os.path.exists(output_path))

        # Verify P5.3.2: Slide count
        prs = Presentation(output_path)
        self.assertEqual(len(prs.slides), 5)

        # Verify P5.3.3: Every slide has text
        for i, slide in enumerate(prs.slides):
            text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
            self.assertGreater(len(text_shapes), 0, f"Slide {i+1} has no text")

        # Verify P5.3.4: Content correctness
        self.assertIn("Project Reef", prs.slides[0].shapes[0].text)
        slide2_text = " ".join(s.text for s in prs.slides[1].shapes if s.has_text_frame)
        self.assertIn("Key Features", slide2_text)
        self.assertIn("Multi-agent", slide2_text)
        self.assertIn("Thank You", prs.slides[4].shapes[0].text)

        # Verify overrides: font_size 48 on slide 5 title
        title_font = prs.slides[4].shapes[0].text_frame.paragraphs[0].runs[0].font.size
        if title_font:
            self.assertEqual(title_font, Pt(48))

        print(f"\nE2E complete: {output_path} ({os.path.getsize(output_path)} bytes, {len(prs.slides)} slides)")

    def test_e2e_cleanup_artifacts_exist(self):
        """P5.3.6: All intermediate artifacts created during pipeline."""
        self.test_e2e_full_pipeline()
        for f in ["template_structure.json", "outline.json", "detail_plan.json",
                   "style_mapping.json", "result.pptx"]:
            path = os.path.join(self.workspace, f)
            self.assertTrue(os.path.exists(path), f"Missing: {f}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
