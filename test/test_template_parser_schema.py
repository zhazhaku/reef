"""
P0.2: Template parser JSON Schema validation.
Validates that the TemplateStructure JSON format contains all required
fields for LLM consumption, using real .pptx templates.
"""
import unittest
import json
import os
from pptx import Presentation
from pptx.util import Pt

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "fixtures", "test_template.pptx")
FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")

# Expected shape type enumeration
VALID_SHAPE_TYPES = {"TEXT_BOX", "PLACEHOLDER", "PICTURE", "TABLE", "CHART", "AUTO_SHAPE", "GROUP"}

# Required top-level keys
REQUIRED_TOP_KEYS = {"slides", "theme"}

# Required slide-level keys
REQUIRED_SLIDE_KEYS = {"slide_number", "slide_type", "shapes"}

# Required shape-level keys
REQUIRED_SHAPE_KEYS = {"type", "left", "top", "width", "height", "name", "shape_id"}


def _mock_parse_template(pptx_path):
    """
    Simulate template_parser output for schema validation.
    This is the contract we will implement in P1.
    """
    prs = Presentation(pptx_path)
    slides = []
    for idx, slide in enumerate(prs.slides):
        shapes = []
        for shape in slide.shapes:
            info = {
                "type": _shape_type_name(shape),
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height,
                "name": shape.name,
                "shape_id": shape.shape_id,
            }
            # Optional text content
            if shape.has_text_frame:
                paragraphs = []
                for p in shape.text_frame.paragraphs:
                    runs_info = []
                    for run in p.runs:
                        runs_info.append({
                            "text": run.text,
                            "font_name": run.font.name,
                            "font_size": _pt_or_none(run.font.size),
                            "bold": run.font.bold,
                            "italic": run.font.italic,
                        })
                    paragraphs.append({
                        "text": p.text,
                        "runs": runs_info,
                        "alignment": str(p.alignment) if p.alignment else None,
                    })
                info["text_content"] = {
                    "full_text": shape.text,
                    "paragraphs": paragraphs,
                }
            shapes.append(info)

        slides.append({
            "slide_number": idx + 1,
            "slide_type": _infer_slide_type(idx, len(prs.slides), slide),
            "shapes": shapes,
        })

    # Theme info (simplified for schema test)
    theme = {
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
    }

    return {"slides": slides, "theme": theme}


def _shape_type_name(shape):
    """Map MSO_SHAPE_TYPE to human-readable name."""
    type_map = {
        "PLACEHOLDER (14)": "PLACEHOLDER",
        "TEXT_BOX (17)": "TEXT_BOX",
        "PICTURE (13)": "PICTURE",
        "TABLE (19)": "TABLE",
        "AUTO_SHAPE (1)": "AUTO_SHAPE",
        "GROUP (6)": "GROUP",
    }
    raw = str(shape.shape_type)
    return type_map.get(raw, raw)


def _pt_or_none(size):
    """Convert EMU font size to pt, or None."""
    if size is None:
        return None
    return size // 12700  # 1 pt = 12700 EMU


def _infer_slide_type(idx, total, slide):
    """Infer slide type from position and shapes."""
    has_title = any(
        s.is_placeholder and s.placeholder_format.idx == 0
        for s in slide.shapes
    )
    if idx == 0 and has_title:
        return "COVER"
    if idx == total - 1:
        return "ENDING"
    return "CONTENT"


class TestTemplateSchemaStructure(unittest.TestCase):
    """P0.2: Schema validation tests."""

    def test_schema_top_level_keys(self):
        """P0.2.2: Output must contain 'slides' and 'theme' top-level keys."""
        result = _mock_parse_template(TEMPLATE_PATH)
        for key in REQUIRED_TOP_KEYS:
            self.assertIn(key, result, f"Missing top-level key: {key}")

    def test_schema_slide_keys(self):
        """Each slide must have slide_number, slide_type, shapes."""
        result = _mock_parse_template(TEMPLATE_PATH)
        for slide in result["slides"]:
            for key in REQUIRED_SLIDE_KEYS:
                self.assertIn(key, slide, f"Missing slide key '{key}' in slide {slide.get('slide_number')}")

    def test_schema_shape_keys(self):
        """P0.2.3: Each shape must have required keys with valid type."""
        result = _mock_parse_template(TEMPLATE_PATH)
        for slide in result["slides"]:
            for shape in slide["shapes"]:
                for key in REQUIRED_SHAPE_KEYS:
                    self.assertIn(key, shape, f"Missing shape key '{key}' in shape '{shape.get('name')}'")
                # type must be valid
                self.assertIn(
                    shape["type"], VALID_SHAPE_TYPES,
                    f"Invalid shape type '{shape['type']}' in shape '{shape['name']}'"
                )

    def test_schema_slide_numbering(self):
        """Slide numbers are sequential starting from 1."""
        result = _mock_parse_template(TEMPLATE_PATH)
        for i, slide in enumerate(result["slides"]):
            self.assertEqual(slide["slide_number"], i + 1)

    def test_schema_cover_slide_has_title(self):
        """COVER slide (slide 1) must have at least one shape with text."""
        result = _mock_parse_template(TEMPLATE_PATH)
        cover = result["slides"][0]
        self.assertEqual(cover["slide_type"], "COVER")
        text_shapes = [s for s in cover["shapes"] if "text_content" in s and s["text_content"]["full_text"]]
        self.assertGreater(len(text_shapes), 0, "COVER slide must have at least one text shape")

    def test_schema_json_serializable(self):
        """Output must be JSON-serializable."""
        result = _mock_parse_template(TEMPLATE_PATH)
        json_str = json.dumps(result, ensure_ascii=False, indent=2)
        # Verify round-trip
        parsed = json.loads(json_str)
        self.assertEqual(len(parsed["slides"]), len(result["slides"]))
        print(f"\nJSON output size: {len(json_str)} bytes, {len(parsed['slides'])} slides")

    def test_schema_multi_template_cross_validation(self):
        """P0.2.4: Multiple templates should all produce valid output."""
        # Create a second template with different structure
        prs2 = Presentation()
        prs2.slide_width = prs2.slide_width
        prs2.slide_height = prs2.slide_height
        prs2.slides.add_slide(prs2.slide_layouts[0])
        prs2.slides.add_slide(prs2.slide_layouts[1])
        path2 = os.path.join(FIXTURES_DIR, "output", "template2.pptx")
        os.makedirs(os.path.dirname(path2), exist_ok=True)
        prs2.save(path2)

        # Test both templates
        for path in [TEMPLATE_PATH, path2]:
            with self.subTest(template=os.path.basename(path)):
                result = _mock_parse_template(path)
                self.assertIn("slides", result)
                self.assertGreater(len(result["slides"]), 0)

        # Cleanup
        os.remove(path2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
